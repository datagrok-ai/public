"""Regenerate data/demo/chem/protac_degraders.csv from PROTAC-PatentDB.

    pip install rdkit pandas openpyxl
    python build.py [out.csv]

Downloads the CC BY 4.0 source (~76 MB), splits every PROTAC into warhead / linker /
E3 ligand (see split.py), keeps the linkers that form a usable matrix, and writes the CSV.
Takes about 6 minutes.
"""
import csv, os, sys, urllib.request
from collections import Counter, defaultdict

import openpyxl
from rdkit import Chem, RDLogger
from rdkit.Chem import Crippen, Descriptors, QED, rdMolDescriptors, rdmolops

from split import split

RDLogger.DisableLog('rdApp.*')

SOURCE = 'https://ndownloader.figshare.com/files/55486868'      # figshare 29351321
XLSX = 'PROTAC_Patent_Compounds.xlsx'
OUT = sys.argv[1] if len(sys.argv) > 1 else 'protac_degraders.csv'

# columns taken from the source sheet
COLS = {'SMILES': 0, 'Target': 1, 'Patent': 4, 'Year': 6, 'Assignee': 7, 'caco2': 43, 'logS': 36}
# VHL PROTACs only ever use the two VH032 variants, so they need a lower column bar
CRITERIA = {'CRBN': (4, 3, 20), 'VHL': (4, 2, 20)}
HDR = ['Compound', 'Warhead', 'Linker', 'E3 Ligand', 'Target', 'E3 Ligase',
       'Average Mass', 'cLogP', 'TPSA', 'HBA', 'HBD', 'Rotatable Bonds', 'Fsp3', 'QED',
       'Caco-2 Permeability (pred)', 'Solubility logS (pred)', 'Patent', 'Year', 'Assignee']


def source_rows():
    if not os.path.exists(XLSX):
        print('downloading', SOURCE)
        urllib.request.urlretrieve(SOURCE, XLSX)
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    it = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    next(it)
    for r in it:
        if r and r[0]:
            yield {k: r[i] for k, i in COLS.items()}
    wb.close()


def zips_back(rec):
    try:
        m = rdmolops.molzip(Chem.MolFromSmiles(rec['linker']), Chem.MolFromSmiles(rec['warhead']))
        return Chem.MolToSmiles(rdmolops.molzip(m, Chem.MolFromSmiles(rec['e3']))) == rec['canonical']
    except Exception:
        return False


def num(v, nd=3):
    try:
        return round(float(v), nd)
    except (TypeError, ValueError):
        return ''


seen, split_ok = set(), []
for n, r in enumerate(source_rows(), 1):
    m = Chem.MolFromSmiles(r['SMILES'])
    if m is None:
        continue
    can = Chem.MolToSmiles(m)
    if can in seen:
        continue
    seen.add(can)
    s = split(r['SMILES'])
    if s:
        r.update(s, canonical=can)
        split_ok.append(r)
    if n % 10000 == 0:
        print(n, 'read,', len(split_ok), 'split', flush=True)
print('read', n, '- split', len(split_ok))

hv = lambda s: Chem.MolFromSmiles(s).GetNumHeavyAtoms()
top_e3 = {e for e, _ in Counter(r['e3'] for r in split_ok).most_common(12)}
rows = [r for r in split_ok if 4 <= hv(r['linker']) <= 24 and r['e3'] in top_e3]
wc = Counter(r['warhead'] for r in rows)
rows = [r for r in rows if wc[r['warhead']] >= 3]

chosen = set()
for fam, (min_w, min_e3, min_rows) in CRITERIA.items():
    g = defaultdict(list)
    for r in rows:
        if r['e3_family'] == fam:
            g[r['linker']].append(r)
    chosen |= {lk for lk, rs in g.items()
               if len({x['warhead'] for x in rs}) >= min_w and len({x['e3'] for x in rs}) >= min_e3
               and len(rs) >= min_rows}

out, bad = [], 0
for r in (x for x in rows if x['linker'] in chosen):
    if not zips_back(r):
        bad += 1
        continue
    m = Chem.MolFromSmiles(r['canonical'])
    out.append([r['canonical'], r['warhead'], r['linker'], r['e3'], r['Target'], r['e3_family'],
                round(Descriptors.MolWt(m), 2), round(Crippen.MolLogP(m), 2),
                round(rdMolDescriptors.CalcTPSA(m), 2), rdMolDescriptors.CalcNumHBA(m),
                rdMolDescriptors.CalcNumHBD(m), rdMolDescriptors.CalcNumRotatableBonds(m),
                round(rdMolDescriptors.CalcFractionCSP3(m), 3), round(QED.qed(m), 3),
                num(r.get('caco2')), num(r.get('logS')), r.get('Patent', ''), r.get('Year', ''),
                r.get('Assignee', '')])

with open(OUT, 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(HDR)
    w.writerows(out)
print('matrices', len(chosen), '- rows', len(out), '- round-trip failures dropped', bad)
